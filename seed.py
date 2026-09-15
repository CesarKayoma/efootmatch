"""
Script para inserir gols em massa na tabela `goals` do eFOOTMATCH.

Lê diretamente um arquivo .xlsx com estas colunas, nesta ordem:

    match_id, scorer_id, assister_id, own_goal, team, created_at, updated_at

- assister_id pode vir vazio (gol sem assistência)
- own_goal: TRUE/FALSE (célula booleana ou texto)
- team: 1 ou 2 (vira team_id)
- created_at / updated_at: data (célula de data ou texto DD/MM/YYYY)

Uso:
    1. Configure DATABASE_URL no ambiente ou num .env
    2. pip install "psycopg[binary]" python-dotenv openpyxl
    3. python seed_goals.py caminho/para/goals.xlsx
"""

import os
import sys
from datetime import datetime

import psycopg
from openpyxl import load_workbook

try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass

EXPECTED_HEADERS = [
    "match_id", "scorer_id", "assister_id", "own_goal", "team", "created_at", "updated_at",
]


def parse_date(value) -> datetime:
    if isinstance(value, datetime):
        return value
    return datetime.strptime(str(value).strip(), "%d/%m/%Y")


def parse_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return str(value).strip().upper() == "TRUE"


def parse_optional_int(value):
    if value is None or str(value).strip() == "":
        return None
    return int(value)


def load_rows(xlsx_path: str):
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    ws = wb.active

    rows_iter = ws.iter_rows(values_only=True)
    header = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
    col_index = {name: header.index(name) for name in EXPECTED_HEADERS if name in header}
    missing = [name for name in EXPECTED_HEADERS if name not in col_index]
    if missing:
        sys.exit(f"Colunas faltando na planilha: {missing}")

    rows = []
    for line_num, row in enumerate(rows_iter, start=2):
        if row is None or all(cell is None for cell in row):
            continue
        try:
            rows.append((
                int(row[col_index["match_id"]]),
                int(row[col_index["scorer_id"]]),
                parse_optional_int(row[col_index["assister_id"]]),
                parse_bool(row[col_index["own_goal"]]),
                int(row[col_index["team"]]),
                parse_date(row[col_index["created_at"]]),
                parse_date(row[col_index["updated_at"]]),
            ))
        except (ValueError, TypeError) as e:
            sys.exit(f"Erro na linha {line_num} da planilha: {row} ({e})")
    return rows


def main():
    if len(sys.argv) != 2:
        sys.exit("Uso: python seed_goals.py caminho/para/goals.xlsx")

    dsn = os.environ.get("DATABASE_URL")
    if not dsn:
        sys.exit("Erro: defina DATABASE_URL no ambiente ou num arquivo .env")

    rows = load_rows(sys.argv[1])

    with psycopg.connect(dsn) as conn:
        with conn.cursor() as cur:
            cur.executemany(
                """
                INSERT INTO goals
                    (match_id, scorer_id, assister_id, own_goal, team_id, created_at, updated_at)
                VALUES (%s, %s, %s, %s, %s, %s, %s)
                """,
                rows,
            )
            inserted = cur.rowcount
        conn.commit()

    print(f"Concluído. Gols inseridos: {inserted}")


if __name__ == "__main__":
    main()